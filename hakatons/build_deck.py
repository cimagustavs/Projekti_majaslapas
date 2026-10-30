"""Builds hakatons/buvgruzu-birza.pptx — the 3-minute pitch deck (import into Canva)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DARK = RGBColor(0x24, 0x4A, 0x30)
MID = RGBColor(0x5E, 0x8C, 0x5A)
LIGHT = RGBColor(0xEE, 0xF3, 0xEA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x26, 0x1E)
MUTED = RGBColor(0x5B, 0x66, 0x5D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xCF, 0xE0, 0xC8)
HEAD = "Montserrat"
BODY = "Open Sans"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def text(slide, x, y, w, h, runs, size=18, color=INK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str or list of paragraphs; a paragraph is str or (str, overrides)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = [runs] if isinstance(runs, str) else runs
    for i, p in enumerate(paras):
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        para.space_after = Pt(o.get("after", 6))
        r = para.add_run()
        r.text = t
        f = r.font
        f.name = o.get("font", font)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return tb


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def placeholder(slide, x, y, w, h, label):
    rect(slide, x, y, w, h, PALE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, x, y, w, h, label.split("\n"), size=16, color=DARK, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)


def kicker(slide, x, y, t, color=MID):
    text(slide, x, y, 8, 0.4, t.upper(), size=13, color=color, font=HEAD, bold=True)


# 1 — Problēma
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
kicker(s, 0.7, 0.6, "Problēma", PALE)
text(s, 0.7, 1.05, 6.6, 2.4, "Būvgruži mežā ir lētāki nekā būvgruži pārstrādē",
     size=40, color=WHITE, font=HEAD, bold=True, spacing=1.0)
text(s, 0.7, 3.55, 6.3, 1.0,
     "Nelegāli izmest būvgružus ir lētāk un vienkāršāk nekā tos sašķirot un nodot legāli.",
     size=18, color=PALE)
for i, (big, small) in enumerate([("~½", "no zemes dzīlēs iegūtajiem resursiem patērē būvniecība"),
                                  ("> ⅓", "no visiem ES atkritumiem ir būvniecības atkritumi")]):
    x = 0.7 + i * 3.25
    text(s, x, 4.75, 3.0, 0.8, big, size=40, color=WHITE, font=HEAD, bold=True)
    text(s, x, 5.55, 2.9, 0.9, small, size=14, color=PALE)
text(s, 0.7, 6.75, 6.5, 0.4, "Avots: hakatona izaicinājums (SCHWENK, RTU, Green Liberty)",
     size=10, color=PALE)
s.shapes.add_picture("hakatons/foto-mezs.jpg", Inches(7.75), Inches(0.6), Inches(4.95), Inches(4.6))
text(s, 7.75, 5.22, 4.95, 0.2, "Foto: iauto.lv", size=8, color=MID, align=PP_ALIGN.RIGHT)
rect(s, 7.75, 5.45, 4.95, 1.45, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.05, 5.6, 4.4, 1.2,
     [("„Izvedam būvgružus lēti!”", {"size": 20, "bold": True, "font": HEAD, "color": DARK}),
      ("Viens zvans — un gruži nonāk mežā.", {"size": 14, "color": MUTED})],
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, "0:35 — Iedomājieties: jūs remontējat dzīvokli, un jums ir 30 maisi ar betonu un flīzēm. "
         "Nav piekabes, nezināt, kur tos nodot. Internetā — \"izvedam būvgružus lēti\". Viens zvans, "
         "un problēma atrisināta. Bet daļa šo gružu nonāk mežā. Būvniecība patērē pusi no visiem zemē "
         "iegūtajiem resursiem, un būvgruži ir vairāk nekā trešdaļa no visiem ES atkritumiem. Problēma "
         "nav cilvēkos. Problēma ir tā, ka nelegālais ceļš ir lētāks un vienkāršāks.\n\n"
         "Foto avots: https://iauto.lv/zinas/latvija/49412-izg-atkritumus-me-k-tdus-cilvkus-nosaukt-foto")

# 2 — Risinājums
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
kicker(s, 0.7, 0.6, "Risinājums")
text(s, 0.7, 1.0, 10, 0.9, "Būvgružu birža", size=44, color=DARK, font=HEAD, bold=True)
text(s, 0.7, 1.85, 10.2, 0.6, "Aprites ekonomikas platforma būvgružiem — savieno cilvēkus, kuriem ir "
     "būvgruži, ar uzņēmumiem, kas tos pārstrādā.", size=18, color=MUTED)
s.shapes.add_picture("hakatons/logo.png", Inches(11.2), Inches(0.55), Inches(1.45))
cards = [("Cilvēkiem", ["Nodod sašķirotus gružus bez maksas",
                        "Par koksni — nauda",
                        "Par betonu un ķieģeļiem — punkti vai pārstrādātas šķembas"]),
         ("Uzņēmumiem", ["Tīrs, sašķirots materiāls",
                         "Bez šķirošanas izmaksām",
                         "Lētāka izejviela nekā jauna"]),
         ("Videi", ["Mazāk izgāztuvju mežos",
                    "Mazāk jaunu izejvielu ieguves",
                    "Pašvaldība ietaupa uz sakopšanu"])]
for i, (h, items) in enumerate(cards):
    x = 0.7 + i * 4.05
    rect(s, x, 3.0, 3.8, 3.8, CARD, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, 3.0, 3.8, 0.12, MID)
    text(s, x + 0.35, 3.4, 3.1, 0.6, h, size=24, color=DARK, font=HEAD, bold=True)
    text(s, x + 0.35, 4.2, 3.1, 2.5, [("— " + t, {"after": 12}) for t in items], size=16)
notes(s, "0:30 — Mūsu risinājums ir Būvgružu birža — aprites ekonomikas platforma, kas savieno cilvēkus, "
         "kuriem ir būvgruži, ar uzņēmumiem, kas tos pārstrādā. Princips kā pudeļu depozītam — nodod un "
         "saņem. Cilvēkiem: bezmaksas nodošana un atlīdzība — par koksni nauda, par betonu punkti. "
         "Uzņēmumiem: tīrs materiāls bez šķirošanas izmaksām. Videi: mazāk izgāztuvju mežos un mazāk "
         "jaunu izejvielu ieguves.")

# 3 — Kā tas darbojas
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
kicker(s, 0.7, 0.6, "Kā tas darbojas")
text(s, 0.7, 1.0, 6.5, 0.9, "Četri soļi lietotnē", size=40, color=DARK, font=HEAD, bold=True)
steps = [("Reģistrējies", "Konts lietotnē"),
         ("Piesaki materiālu", "Foto, veids, kvalitāte, daudzums"),
         ("Izvēlies punktu kartē", "Tikai tur, kur atlīdzība > degvielas izmaksas"),
         ("Nodod un saņem", "Nosver, pārbauda, atlīdzība kontā")]
for i, (h, sub) in enumerate(steps):
    y = 2.2 + i * 1.2
    c = rect(s, 0.7, y, 0.8, 0.8, DARK, MSO_SHAPE.OVAL)
    text(s, 0.7, y, 0.8, 0.8, str(i + 1), size=24, color=WHITE, font=HEAD, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.8, y - 0.02, 5.2, 0.5, h, size=22, color=INK, font=HEAD, bold=True)
    text(s, 1.8, y + 0.45, 5.2, 0.4, sub, size=15, color=MUTED)
rect(s, 8.3, 0.6, 3.6, 6.3, DARK, MSO_SHAPE.ROUNDED_RECTANGLE)
placeholder(s, 8.45, 0.75, 3.3, 6.0, "VIDEO\nno lietotnes\n(ievelc Canva)")
notes(s, "0:15 — Lietotnē tie ir četri soļi: reģistrējies, piesaki materiālu ar foto, izvēlies punktu "
         "kartē un saņem atlīdzību. Parādīšu, kā tas izskatās.\n\n[VIDEO 1:00 — Gustavs komentē]")

# 4 — Kāpēc tieši mūsu risinājums
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
kicker(s, 0.7, 0.6, "Kāpēc tieši mūsu risinājums")
text(s, 0.7, 1.0, 12, 0.9, "Izdevīgi, nevis tikai pareizi", size=40, color=DARK, font=HEAD, bold=True)
reasons = [("Izdevīgi", "Nepaļaujamies uz apziņu — legālais ceļš kļūst lētāks par mežu"),
           ("Reāli aprēķināts", "Atlīdzība balstīta uz pārstrādātāju cenām un peļņu (~2 %)"),
           ("Atmaksājas braukt", "Lietotne rāda tikai punktus, kur atlīdzība > degvielas izmaksas"),
           ("Bez jaunas infrastruktūras", "Izmanto esošos licencētos punktus — sākam ar vienu pašvaldību")]
for i, (h, sub) in enumerate(reasons):
    x = 0.7 + (i % 2) * 6.05
    y = 2.15 + (i // 2) * 1.6
    rect(s, x, y, 5.85, 1.4, CARD, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x + 0.3, y + 0.3, 0.8, 0.8, DARK, MSO_SHAPE.OVAL)
    text(s, x + 0.3, y + 0.3, 0.8, 0.8, str(i + 1), size=22, color=WHITE, font=HEAD, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.35, y + 0.22, 4.3, 0.5, h, size=20, color=INK, font=HEAD, bold=True)
    text(s, x + 1.35, y + 0.7, 4.3, 0.6, sub, size=14, color=MUTED)
rect(s, 0.7, 5.45, 11.9, 1.1, DARK, MSO_SHAPE.ROUNDED_RECTANGLE)
prices = [("Koksne", "15–30 €/t"), ("Betons", "1–3 €/t"), ("Ķieģeļi", "0,5–1,5 €/t")]
for i, (m, p) in enumerate(prices):
    x = 1.0 + i * 2.75
    text(s, x, 5.55, 2.6, 0.4, m, size=13, color=PALE)
    text(s, x, 5.9, 2.6, 0.5, p, size=22, color=WHITE, font=HEAD, bold=True)
text(s, 9.15, 5.55, 3.3, 0.9, [("Atlīdzība =", {"size": 13, "color": PALE}),
                               ("svars × cena × kvalitāte", {"size": 15, "bold": True, "font": HEAD})],
     color=WHITE)
text(s, 0.7, 6.7, 11.9, 0.3, "€/t — aplēses pēc uzņēmumu datiem, jāpārbauda pilotprojektā.",
     size=10, color=MUTED)
notes(s, "0:30 — Kāpēc tieši mūsu risinājums? Pirmkārt, tas ir izdevīgi, nevis tikai pareizi — mēs "
         "nepaļaujamies uz cilvēku apziņu. Otrkārt, mēs to aprēķinājām: pārstrādātāju peļņa ir ap diviem "
         "procentiem, tāpēc par koksni maksājam naudu, par betonu — punktus. Treškārt, lietotne rāda tikai "
         "tos punktus, kur braukt atmaksājas. Un mums nevajag jaunu infrastruktūru — sākam ar "
         "pilotprojektu vienā pašvaldībā.")

# 5 — Noslēgums
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
s.shapes.add_picture("hakatons/logo.png", Inches(5.92), Inches(0.7), Inches(1.5))
text(s, 0.7, 2.5, 11.9, 2.0, "Lai būvgruži atgrieztos būvniecībā, nevis mežā",
     size=44, color=WHITE, font=HEAD, bold=True, align=PP_ALIGN.CENTER, spacing=1.0)
text(s, 0.7, 4.55, 11.9, 0.7, "Būvgružu birža", size=28, color=PALE, font=HEAD, bold=True,
     align=PP_ALIGN.CENTER)
text(s, 0.7, 5.2, 11.9, 0.5, "Komanda Alanguka", size=18, color=PALE, align=PP_ALIGN.CENTER)
text(s, 0.7, 6.3, 11.9, 0.6, "Paldies! Jautājumi?", size=22, color=WHITE, align=PP_ALIGN.CENTER)
notes(s, "0:10 — Būvgružu birža — lai būvgruži atgrieztos būvniecībā, nevis mežā. Paldies! "
         "Labprāt atbildēsim uz jautājumiem.")

prs.save("hakatons/buvgruzu-birza.pptx")
